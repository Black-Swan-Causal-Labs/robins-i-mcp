"""Ingestion: document -> SectionMap with character-offset spans.

Ported from the TARGET server's ingestion layer, which is guideline-agnostic.
The SectionMap is the deterministic substrate every assessment runs over: an
evidence span is an offset into `full_text` as produced here, so the map carries
an extractor stamp and a sha256 of the normalized text. A span is only
meaningful alongside those two values.

Two things are new for ROBINS-I:

* **Cue detection is per-domain.** TARGET only needed to know whether a
  target-trial protocol table and a flow diagram were present. ROBINS-I answers
  depend on whether the paper discusses negative controls (1.4), blinding
  (5.2), missing data (domain 4), a registered analysis plan (6.1), and so on.
  `CUES` maps those to search terms, and `detect_cues` reports which fired.

* **Absence is evidence, so absence has to be recorded.** A `manuscript_absent`
  answer asserts something is not in the paper. `search_record` runs the terms
  over named sections and returns a structured, reproducible statement of what
  was searched and what was found — so an NI or N answer can be audited instead
  of taken on trust.
"""

from __future__ import annotations

import functools
import hashlib
import re
import unicodedata
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any, Iterable, Sequence

from pypdf import PdfReader

INGEST_VERSION = "0.1.0"


def _extractor_stamp(engine: str) -> str:
    return f"robins-i-mcp-ingest/{INGEST_VERSION} ({engine})"


EXTRACTOR_VERSION = _extractor_stamp("pypdf")


class ExtractionError(ValueError):
    """A document yielded implausibly little text (e.g. a scanned PDF with no
    text layer). Raised so callers fail loudly instead of silently assessing
    near-empty content."""


CANONICAL_SECTIONS = ("abstract", "introduction", "methods", "results", "discussion", "other")

_HEADING_PATTERNS: list[tuple[str, re.Pattern[str]]] = [
    ("abstract", re.compile(r"^\s*abstract\s*$", re.IGNORECASE)),
    ("introduction", re.compile(r"^\s*(?:\d+\.?\s*)?(introduction|background)\s*$", re.IGNORECASE)),
    ("methods", re.compile(
        r"^\s*(?:\d+\.?\s*)?(methods?|materials and methods|patients and methods"
        r"|study design and methods)\s*$", re.IGNORECASE)),
    ("results", re.compile(r"^\s*(?:\d+\.?\s*)?results?\s*$", re.IGNORECASE)),
    ("discussion", re.compile(r"^\s*(?:\d+\.?\s*)?(discussion|comment)\s*$", re.IGNORECASE)),
    ("other", re.compile(
        r"^\s*(?:\d+\.?\s*)?(references|acknowledg(e)?ments?|funding|declarations"
        r"|supplementary (material|information))\s*$", re.IGNORECASE)),
]

SUPPLEMENT_STATES = ("retrieved", "user_provided", "none_exists", "not_retrieved", "not_checked")


# --------------------------------------------------------------------------- #
# Cues — where each domain's evidence tends to live, and what to search for
# when an assessor wants to assert absence.
# --------------------------------------------------------------------------- #

@dataclass(frozen=True)
class Cue:
    key: str
    #: signalling questions this cue informs
    questions: tuple[str, ...]
    terms: tuple[str, ...]
    #: sections worth searching; empty means the whole bundle
    sections: tuple[str, ...] = ()
    note: str = ""


CUES: tuple[Cue, ...] = (
    Cue("target_trial_table", ("C1", "C2", "C3", "C4"),
        ("target trial", "protocol component", "specification", "emulation"),
        note="C1-C3 usually come from a specification/emulation table, often in the supplement."),
    Cue("estimand", ("C4",),
        ("intention-to-treat", "intention to treat", "per-protocol", "per protocol",
         "causal contrast", "estimand"),
        note="C4 selects the domain 1 variant, so this cue gates half the assessment."),
    Cue("deviation_handling", ("C4", "1.1B"),
        ("censor", "switching", "switch", "discontinu", "adherence", "protocol deviation",
         "inverse probability of censoring", "g-formula", "g-method", "clone"),
        note="Censoring or follow-up partitioning at deviation is what makes C4 = yes."),
    Cue("confounder_control", ("1.1", "1.2", "1.3"),
        ("adjust", "matched", "matching", "propensity", "stratif", "inverse probability",
         "standardiz", "covariate"),
        sections=("methods",)),
    Cue("negative_control", ("1.4", "1.5B"),
        ("negative control", "falsification", "quantitative bias analysis", "E-value",
         "bias analysis", "unmeasured confounding"),
        note="Answering 1.4 N without searching for these is unsafe."),
    Cue("exposure_ascertainment", ("2.1", "2.4", "2.5"),
        ("exposure", "classif", "dispens", "prescription", "new user", "new-user",
         "index date", "misclassif"),
        sections=("methods",)),
    Cue("followup_start", ("3.1", "3.2"),
        ("follow-up", "follow up", "baseline", "index date", "time zero", "landmark",
         "immortal", "prevalent user", "washout", "run-in", "blanking"),
        note="Immortal time and prevalent-user problems live here."),
    Cue("selection", ("3.3", "3.4", "3.5"),
        ("eligib", "inclusion", "exclusion", "excluded", "restrict", "selection"),
        ),
    Cue("selection_correction", ("3.6", "3.7", "3.8"),
        ("inverse probability of selection", "selection weight", "sensitivity analys",
         "bias analysis"),
        ),
    Cue("missing_data", ("4.1", "4.2", "4.3", "4.4"),
        ("missing", "complete case", "complete-case", "loss to follow-up",
         "lost to follow-up", "not available", "unavailable", "did not have",
         "not recorded", "unrecorded", "no recorded", "incomplete"),
        ),
    Cue("imputation", ("4.7", "4.8", "4.9", "4.10"),
        ("imput", "missing at random", "MAR", "MCAR", "MNAR", "chained equations",
         "full information maximum likelihood", "inverse probability weight"),
        ),
    Cue("outcome_ascertainment", ("5.1",),
        ("outcome", "ascertain", "diagnos", "registry", "linkage", "surveillance",
         "detection"),
        sections=("methods",)),
    Cue("blinding", ("5.2", "5.3"),
        ("blind", "masked", "unaware", "adjudicat", "independent review"),
        note="Observational studies rarely blind; 5.2 is usually answered from absence."),
    Cue("prespecification", ("6.1",),
        ("protocol", "prespecified", "pre-specified", "analysis plan",
         "statistical analysis plan", "registered", "registration", "ClinicalTrials.gov",
         "EUPAS", "EU PAS", "OSF"),
        note="6.1 is answerable only against a plan; absence here is the usual finding."),
    Cue("multiplicity", ("6.2", "6.3", "6.4"),
        ("subgroup", "sensitivity analys", "secondary analys", "alternative",
         "additional analys", "post hoc", "post-hoc"),
        ),
)

CUES_BY_KEY = {c.key: c for c in CUES}


@functools.lru_cache(maxsize=512)
def _term_pattern(term: str) -> re.Pattern[str]:
    """Compile a search term.

    Acronyms and very short terms are anchored to word boundaries; anything
    longer is matched as a substring so that deliberate stems ("imput",
    "classif", "adjudicat") catch their whole family. Without the anchor,
    "MAR" matches inside "marked" and a paper that never imputed anything
    reports dozens of imputation hits.
    """
    escaped = re.escape(term)
    if term.isupper() or len(term) <= 4:
        escaped = rf"\b{escaped}\b"
    return re.compile(escaped, re.IGNORECASE)


# --------------------------------------------------------------------------- #
# Data model
# --------------------------------------------------------------------------- #

@dataclass
class Section:
    name: str
    heading: str
    start: int
    end: int
    source: str = "main"


@dataclass
class Hit:
    """One occurrence of a search term."""
    term: str
    start: int
    end: int
    section: str
    source: str
    excerpt: str


@dataclass
class SearchRecord:
    """A reproducible statement of what was searched and what was found.

    This is what a `manuscript_absent` answer must carry: not a prose claim that
    something is missing, but the terms used, the sections covered, and the hit
    count — so a reader can tell an exhaustive search from a cursory one.
    """
    terms: tuple[str, ...]
    sections_searched: tuple[str, ...]
    n_hits: int
    hits: tuple[Hit, ...] = ()
    cue: str = ""

    @property
    def is_absent(self) -> bool:
        return self.n_hits == 0

    def describe(self) -> str:
        where = ", ".join(self.sections_searched) or "whole bundle"
        return (
            f"Searched {where} for {len(self.terms)} term(s) "
            f"({'; '.join(self.terms)}) — {self.n_hits} hit(s)."
        )


@dataclass
class EvidenceSpan:
    """A quote resolved to character offsets in the bundle."""
    quote: str
    start: int
    end: int
    section: str
    source: str
    #: which matching pass succeeded — see MATCH_PASSES. Anything other than
    #: "exact" means the quote needed normalization to land, which is worth
    #: showing rather than hiding.
    match: str = "exact"

    @property
    def locator(self) -> str:
        where = self.section.replace("_", " ")
        if self.source.startswith("supplement:"):
            return f"{where} — {self.source.split(':', 1)[1]}"
        return f"{where} — main text"


class QuoteNotFound(ValueError):
    """A quote could not be resolved to a span in the ingested bundle."""

    def __init__(self, quote: str, question: str = "", near_miss: str = "") -> None:
        self.quote = quote
        self.question = question
        self.near_miss = near_miss
        where = f"{question}: " if question else ""
        hint = f" {near_miss}" if near_miss else ""
        super().__init__(
            f"{where}quote could not be located in the ingested document "
            f"(first 80 chars: {quote[:80]!r}). Evidence must be copied verbatim "
            f"from the parsed bundle.{hint}"
        )


@dataclass
class SectionMap:
    source: str
    manuscript_id: str
    extractor_version: str
    text_sha256: str
    full_text: str
    sections: list[Section] = field(default_factory=list)
    n_pages: int | None = None
    warnings: list[str] = field(default_factory=list)
    supplement_status: str = "not_checked"
    documents: list[dict[str, Any]] = field(default_factory=list)
    citation: str = ""

    # -- lookup ----------------------------------------------------------- #

    def section_text(self, name: str) -> str:
        return "\n\n".join(
            self.full_text[s.start:s.end] for s in self.sections if s.name == name
        )

    def locate(self, quote: str) -> tuple[int, int, str] | None:
        """Resolve a verbatim-ish quote to (start, end, match_pass) in full_text.

        Tries progressively more forgiving projections and stops at the first
        that hits, so an exact match is never reported as a loose one. All are
        whitespace-insensitive and case-insensitive, because PDF extraction
        routinely mangles both at line and column breaks.
        """
        for level in MATCH_PASSES:
            proj_text, index_map = _reduce(self.full_text, level)
            proj_quote, _ = _reduce(quote, level)
            if not proj_quote:
                return None
            pos = proj_text.lower().find(proj_quote.lower())
            if pos >= 0:
                return (index_map[pos], index_map[pos + len(proj_quote) - 1] + 1, level)
        return None

    def resolve(self, quote: str, question: str = "") -> EvidenceSpan:
        found = self.locate(quote)
        if found is None:
            raise QuoteNotFound(quote, question, self._near_miss(quote))
        start, end, level = found
        return EvidenceSpan(
            quote=quote, start=start, end=end, match=level,
            section=self.section_at(start), source=self.source_at(start),
        )

    def _near_miss(self, quote: str, probe_words: int = 6) -> str:
        """Best-effort locator for a quote that did not resolve, so the failure
        says where the assessor probably meant rather than only that it failed.

        Slides the probe window through the quote rather than only testing its
        prefix. The common cause of a near miss is a quote spanning a page break
        with a table float or running header interposed, in which case the
        opening words are contiguous with nothing and only a later window hits.
        """
        # Word-level probing needs whitespace, so this uses the "exact"
        # projection rather than the looser passes, which strip spaces entirely.
        proj_text, index_map = _reduce(self.full_text, "exact")
        lowered = proj_text.lower()
        words = _reduce(quote, "exact")[0].split()
        for n in range(min(probe_words, len(words)), 2, -1):
            for start in range(0, len(words) - n + 1):
                probe = " ".join(words[start:start + n])
                pos = lowered.find(probe.lower())
                if pos < 0:
                    continue
                orig = index_map[pos]
                actual = " ".join(self.full_text[orig:orig + 220].split())
                where = "" if start == 0 else f" (matched from word {start + 1} of the quote)"
                return (
                    f"nearest text at offset {orig} ({self.section_at(orig)}){where}: "
                    f"{actual!r}"
                )
        return "no overlapping text found — the quote may be from another document"

    def section_at(self, offset: int) -> str:
        for s in self.sections:
            if s.start <= offset < s.end:
                return s.name
        return "other"

    def source_at(self, offset: int) -> str:
        for s in self.sections:
            if s.start <= offset < s.end:
                return s.source
        return "main"

    # -- search ----------------------------------------------------------- #

    def _search_ranges(self, sections: Sequence[str]) -> list[tuple[int, int]]:
        if not sections:
            return [(0, len(self.full_text))]
        ranges = [(s.start, s.end) for s in self.sections if s.name in sections]
        # Supplements are always in scope: methods detail routinely lives there.
        ranges += [(s.start, s.end) for s in self.sections if s.name == "supplement"]
        return ranges or [(0, len(self.full_text))]

    def search(
        self,
        terms: Iterable[str],
        sections: Sequence[str] = (),
        *,
        max_hits: int = 12,
        cue: str = "",
    ) -> SearchRecord:
        """Case-insensitive search for any of `terms`, restricted to `sections`."""
        terms = tuple(terms)
        ranges = self._search_ranges(sections)
        hits: list[Hit] = []
        for term in terms:
            pattern = _term_pattern(term)
            for start, end in ranges:
                for m in pattern.finditer(self.full_text, start, end):
                    excerpt = ""
                    if len(hits) < max_hits:
                        excerpt = " ".join(
                            self.full_text[max(0, m.start() - 90):m.end() + 90].split()
                        )
                    hits.append(Hit(term, m.start(), m.end(),
                                    self.section_at(m.start()),
                                    self.source_at(m.start()), excerpt))
        searched = tuple(sections) if sections else ("whole bundle",)
        return SearchRecord(
            terms=terms, sections_searched=searched,
            n_hits=len(hits), hits=tuple(h for h in hits if h.excerpt), cue=cue,
        )

    def search_cue(self, key: str, *, max_hits: int = 12) -> SearchRecord:
        """Run the registered search for a named cue."""
        cue = CUES_BY_KEY[key]
        return self.search(cue.terms, cue.sections, max_hits=max_hits, cue=key)

    def detect_cues(self) -> dict[str, SearchRecord]:
        """Run every cue. The result tells an assessor where to look, and gives
        the evidential basis for any answer that rests on absence."""
        return {c.key: self.search_cue(c.key) for c in CUES}

    # -- serialization ---------------------------------------------------- #

    def to_dict(self, include_text: bool = False) -> dict[str, Any]:
        d = asdict(self)
        if not include_text:
            d["full_text"] = f"<{len(self.full_text)} chars omitted>"
        return d


# --------------------------------------------------------------------------- #
# Text handling
# --------------------------------------------------------------------------- #

def _collapse_ws(text: str) -> tuple[str, list[int]]:
    """Collapse whitespace runs to single spaces; return collapsed text and a
    map from collapsed index -> original index."""
    out: list[str] = []
    index_map: list[int] = []
    in_ws = True
    for i, ch in enumerate(text):
        if ch.isspace():
            if not in_ws:
                out.append(" ")
                index_map.append(i)
                in_ws = True
        else:
            out.append(ch)
            index_map.append(i)
            in_ws = False
    if out and out[-1] == " ":
        out.pop()
        index_map.pop()
    return "".join(out), index_map


def _normalize(text: str) -> str:
    text = unicodedata.normalize("NFKC", text)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    # De-hyphenate line breaks. Publisher PDFs emit both "confound-\ning" and
    # "confound -\ning" (a space before the hyphen), so allow optional spaces on
    # either side of the break.
    text = re.sub(r"(\w)[ \t]*-[ \t]*\n[ \t]*(\w)", r"\1\2", text)
    return text


# Quote matching runs over progressively more forgiving projections of the text.
# `full_text` itself is never altered — these are match-time reductions only, and
# the pass that succeeded is recorded on the resulting span so a loose match is
# visible in the report rather than silently equated with an exact one.
#: Progressively more forgiving, and cumulative — each pass relaxes everything
#: the previous one did plus one more thing.
#:   exact          collapse whitespace runs, ignore case
#:   hyphen_relaxed additionally ignore hyphens and all whitespace, so
#:                  "SARS-CoV-2", "SARSCoV-2" and "SARS CoV 2" are one string
#:   refs_stripped  additionally drop inline superscript reference numerals
MATCH_PASSES = ("exact", "hyphen_relaxed", "refs_stripped")

_REF_TOKEN_RE = re.compile(r"(?<=[A-Za-z,.;:)])\s\d{1,3}(?=\s|[,.;:])")


def _reduce(text: str, level: str) -> tuple[str, list[int]]:
    """Project text for matching; return (projected, map to original offsets)."""
    drop = [False] * len(text)

    if level == "refs_stripped":
        for m in _REF_TOKEN_RE.finditer(text):
            for i in range(m.start(), m.end()):
                drop[i] = True

    out: list[str] = []
    index_map: list[int] = []

    if level == "exact":
        in_ws = True
        for i, ch in enumerate(text):
            if ch.isspace():
                if not in_ws:
                    out.append(" ")
                    index_map.append(i)
                    in_ws = True
            else:
                out.append(ch)
                index_map.append(i)
                in_ws = False
        if out and out[-1] == " ":
            out.pop()
            index_map.pop()
        return "".join(out), index_map

    # hyphen_relaxed / refs_stripped: whitespace and hyphens carry no signal.
    for i, ch in enumerate(text):
        if drop[i] or ch.isspace() or ch in "-‐‑‒–—":
            continue
        out.append(ch)
        index_map.append(i)
    return "".join(out), index_map


def _pypdf_pages(path: Path) -> list[str]:
    return [p.extract_text() or "" for p in PdfReader(str(path)).pages]


def _pdfplumber_pages(path: Path) -> list[str] | None:
    """Second-chance extractor: some publisher PDFs have a weak text layer under
    pypdf that pdfplumber recovers."""
    try:
        import pdfplumber
    except ImportError:
        return None
    try:
        with pdfplumber.open(str(path)) as pdf:
            return [(pg.extract_text() or "") for pg in pdf.pages]
    except Exception:
        return None


def _too_sparse(pages: list[str]) -> bool:
    chars = len("".join(pages).strip())
    return len(pages) >= 1 and (chars < 100 or chars < 25 * len(pages))


def _extract_pdf(path: Path) -> tuple[list[str], str]:
    pages, engine = _pypdf_pages(path), "pypdf"
    if _too_sparse(pages):
        alt = _pdfplumber_pages(path)
        if alt is not None and len("".join(alt).strip()) > len("".join(pages).strip()):
            pages, engine = alt, "pdfplumber"
    return pages, engine


def _assert_plausible(text: str, n_pages: int | None, source: str) -> None:
    if n_pages and len(text.strip()) < 100:
        raise ExtractionError(
            f"{source}: extracted only {len(text.strip())} characters from {n_pages} "
            "page(s) — the PDF likely has no usable text layer (scanned or watermarked) "
            "and needs OCR. Not proceeding, to avoid assessing near-empty content."
        )


# --------------------------------------------------------------------------- #
# Entry points
# --------------------------------------------------------------------------- #

def parse_pdf(path: str | Path, manuscript_id: str | None = None) -> SectionMap:
    path = Path(path)
    pages, engine = _extract_pdf(path)
    full_text = _normalize("\n".join(pages))
    _assert_plausible(full_text, len(pages), str(path))
    return _build_map(full_text, source=str(path), manuscript_id=manuscript_id or path.stem,
                      n_pages=len(pages), extractor_version=_extractor_stamp(engine))


def parse_text(text: str, source: str = "<text>",
               manuscript_id: str = "manuscript") -> SectionMap:
    return _build_map(_normalize(text), source=source,
                      manuscript_id=manuscript_id, n_pages=None)


def extract_file(path: str | Path) -> tuple[str, int | None]:
    """Extract raw (un-normalized) text from a file by extension."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        pages, _ = _extract_pdf(path)
        text = "\n".join(pages)
        _assert_plausible(text, len(pages), str(path))
        return text, len(pages)
    if suffix == ".docx":
        try:
            import docx
        except ImportError as e:
            raise RuntimeError("Reading .docx supplements requires python-docx") from e
        d = docx.Document(str(path))
        blocks = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                blocks.append("\t".join(c.text for c in row.cells))
        return "\n".join(blocks), None
    return path.read_text(encoding="utf-8", errors="replace"), None


_DOC_EXTENSIONS = (".pdf", ".docx", ".doc", ".txt", ".md", ".rtf")


def _looks_like_path(s: str) -> bool:
    s = s.strip()
    if not s or "\n" in s or len(s) > 400:
        return False
    return s.lower().endswith(_DOC_EXTENSIONS) or "/" in s or "\\" in s


def parse_document(path_or_text: str, manuscript_id: str | None = None) -> SectionMap:
    """Dispatch: an existing file path -> parse the file; raw text -> parse_text.

    A string that looks like a path but does not resolve raises loudly rather
    than being ingested as its own 30-character 'manuscript'."""
    p = Path(path_or_text)
    try:
        is_file = p.is_file()
    except OSError:  # raw text can exceed max path component length
        is_file = False
    if is_file:
        if p.suffix.lower() == ".pdf":
            return parse_pdf(p, manuscript_id)
        text, _ = extract_file(p)
        return parse_text(text, source=str(p), manuscript_id=manuscript_id or p.stem)
    if _looks_like_path(path_or_text):
        import os
        raise FileNotFoundError(
            f"{path_or_text!r} looks like a file path but no such file exists "
            f"(server working directory: {os.getcwd()!r}). Pass an ABSOLUTE path, "
            "or pass the document text itself."
        )
    return parse_text(path_or_text, manuscript_id=manuscript_id or "manuscript")


def _build_map(full_text: str, source: str, manuscript_id: str, n_pages: int | None,
               extractor_version: str = EXTRACTOR_VERSION) -> SectionMap:
    warnings: list[str] = []
    boundaries: list[tuple[int, str, str]] = []

    offset = 0
    for line in full_text.split("\n"):
        stripped = line.strip()
        if stripped and len(stripped) <= 60:
            for canonical, pat in _HEADING_PATTERNS:
                if pat.match(stripped):
                    boundaries.append((offset, canonical, stripped))
                    break
        offset += len(line) + 1

    boundaries, abstract_warn = _drop_structured_abstract_labels(boundaries, len(full_text))
    warnings.extend(abstract_warn)

    seen: dict[str, int] = {}
    ordered: list[tuple[int, str, str]] = []
    rank = {name: i for i, name in enumerate(CANONICAL_SECTIONS)}
    last_rank = -1
    for off, canonical, heading in boundaries:
        if canonical in seen:
            continue
        if rank[canonical] < last_rank:
            warnings.append(f"Out-of-order heading {heading!r} at {off} ignored")
            continue
        seen[canonical] = off
        ordered.append((off, canonical, heading))
        last_rank = rank[canonical]

    sections: list[Section] = []
    if not ordered:
        warnings.append("No section headings detected; whole document mapped as 'other'")
        sections.append(Section("other", "", 0, len(full_text)))
    else:
        if ordered[0][0] > 0:
            name = "abstract" if "abstract" not in seen else "other"
            sections.append(Section(name, "", 0, ordered[0][0]))
            if name == "abstract":
                warnings.append("No 'Abstract' heading; front matter mapped as abstract")
        for i, (off, canonical, heading) in enumerate(ordered):
            end = ordered[i + 1][0] if i + 1 < len(ordered) else len(full_text)
            sections.append(Section(canonical, heading, off, end))

    missing = [s for s in ("abstract", "methods", "results")
               if s not in {x.name for x in sections}]
    if missing:
        warnings.append(f"Sections not detected: {missing}")

    digest = hashlib.sha256(full_text.encode("utf-8")).hexdigest()
    return SectionMap(
        source=source, manuscript_id=manuscript_id, extractor_version=extractor_version,
        text_sha256=digest, full_text=full_text, sections=sections, n_pages=n_pages,
        warnings=warnings, supplement_status="not_checked",
        documents=[{
            "source": "main", "kind": "main", "filename": source,
            "char_start": 0, "char_end": len(full_text),
            "sha256": digest, "n_pages": n_pages,
        }],
    )


#: A structured abstract (NEJM, BMJ, JAMA and most clinical journals) carries
#: its own BACKGROUND / METHODS / RESULTS / CONCLUSIONS labels on page one.
#: Matched naively, those labels become the body section boundaries — which is
#: how a Methods section ends up 800 characters long while the real one is
#: mapped as Results. Detect the cluster and drop it.
_ABSTRACT_LABELS = ("introduction", "methods", "results", "discussion")


def _is_shouted(heading: str) -> bool:
    """True for a heading rendered in capitals — how clinical journals set the
    labels inside a structured abstract."""
    letters = [c for c in heading if c.isalpha()]
    return bool(letters) and all(c.isupper() for c in letters)


def _drop_structured_abstract_labels(
    boundaries: list[tuple[int, str, str]], text_len: int
) -> tuple[list[tuple[int, str, str]], list[str]]:
    """Remove page-one structured-abstract labels from the boundary list.

    The reliable signal is typographic, not positional: NEJM, JAMA and BMJ set
    the abstract's BACKGROUND / METHODS / RESULTS labels in capitals, while the
    body headings are title case. So a capitalized heading is dropped when the
    same section kind reappears later in mixed case.

    Distance was tried first and is not dependable — a short paper packs its
    body headings as tightly as a long one packs its abstract labels.
    """
    dropped_kinds: set[str] = set()
    cutoff = 0
    for off, canonical, heading in boundaries:
        if canonical not in _ABSTRACT_LABELS or not _is_shouted(heading):
            continue
        repeats_later = any(
            later_off > off and later_kind == canonical and not _is_shouted(later_heading)
            for later_off, later_kind, later_heading in boundaries
        )
        if repeats_later:
            dropped_kinds.add(canonical)
            cutoff = max(cutoff, off)

    if not dropped_kinds:
        return boundaries, []

    kept = [b for b in boundaries
            if not (b[0] <= cutoff and b[1] in dropped_kinds and _is_shouted(b[2]))]
    labels = ", ".join(sorted(dropped_kinds))
    return kept, [
        f"Structured-abstract labels ({labels}) detected before offset {cutoff} "
        "and excluded from body sectioning"
    ]


_SUPPLEMENT_SEP = "\n\n===== SUPPLEMENTARY MATERIAL: {name} =====\n\n"


def build_bundle(main: SectionMap, supplements: list[tuple[str, str, int | None]],
                 supplement_status: str) -> SectionMap:
    """Merge a main-text map with supplement documents into one bundle.

    Supplements matter more for ROBINS-I than for a reporting checklist: the
    target-trial specification that settles C1-C4, and the missing-data and
    analysis detail that domains 1 and 4 turn on, routinely live only in the
    appendix. A bundle assembled without them will read `NI` where the answer
    was simply in a file nobody passed.
    """
    if supplement_status not in SUPPLEMENT_STATES:
        raise ValueError(f"Unknown supplement_status {supplement_status!r}")
    if not supplements:
        main.supplement_status = supplement_status
        return main

    parts = [main.full_text]
    sections = [Section(s.name, s.heading, s.start, s.end, "main") for s in main.sections]
    documents = list(main.documents)
    cursor = len(main.full_text)
    for filename, text, n_pages in supplements:
        sep = _SUPPLEMENT_SEP.format(name=filename)
        text = _normalize(text)
        parts.append(sep)
        parts.append(text)
        seg_start = cursor + len(sep)
        seg_end = seg_start + len(text)
        src = f"supplement:{filename}"
        sections.append(Section("supplement", f"SUPPLEMENT: {filename}", seg_start, seg_end, src))
        documents.append({
            "source": src, "kind": "supplement", "filename": filename,
            "char_start": seg_start, "char_end": seg_end,
            "sha256": hashlib.sha256(text.encode("utf-8")).hexdigest(),
            "n_pages": n_pages,
        })
        cursor = seg_end

    full_text = "".join(parts)
    return SectionMap(
        source=main.source, manuscript_id=main.manuscript_id,
        extractor_version=EXTRACTOR_VERSION,
        text_sha256=hashlib.sha256(full_text.encode("utf-8")).hexdigest(),
        full_text=full_text, sections=sections, n_pages=main.n_pages,
        warnings=list(main.warnings), supplement_status=supplement_status,
        documents=documents, citation=main.citation,
    )
